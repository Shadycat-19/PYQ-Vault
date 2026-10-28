import os
import io
import uuid
import zipfile
import textwrap
from datetime import datetime
from urllib.parse import unquote
from flask import (
    Flask, render_template, request, redirect, url_for, flash, 
    jsonify, send_file, abort
)
from werkzeug.utils import secure_filename
from PIL import Image, ImageDraw, ImageFont
import pypdfium2 as pdfium
import docx

from config import Config
from models import db, Paper, PaperPage, Solution, SolutionPage

def create_app(test_config=None):
    app = Flask(__name__)
    app.config.from_object(Config)
    if test_config:
        app.config.update(test_config)

    db.init_app(app)

    os.makedirs(app.config['PAPERS_UPLOAD_DIR'], exist_ok=True)
    os.makedirs(app.config['SOLUTIONS_UPLOAD_DIR'], exist_ok=True)

    with app.app_context():
        db.create_all()
        for col, col_type in [
            ('text_content', 'TEXT'),
            ('solution_type', 'VARCHAR(20) DEFAULT "images"'),
            ('question_number', 'VARCHAR(50) DEFAULT "General"')
        ]:
            try:
                db.session.execute(db.text(f"ALTER TABLE solution ADD COLUMN {col} {col_type}"))
                db.session.commit()
            except Exception:
                db.session.rollback()

        for col, col_type in [
            ('original_file_path', 'VARCHAR(300)'),
            ('original_filename', 'VARCHAR(200)'),
            ('file_type', 'VARCHAR(20) DEFAULT "images"')
        ]:
            try:
                db.session.execute(db.text(f"ALTER TABLE paper ADD COLUMN {col} {col_type}"))
                db.session.commit()
            except Exception:
                db.session.rollback()

    def render_text_to_page_images(lines, title="Question Paper"):
        font_path = '/System/Library/Fonts/Helvetica.ttc'
        if not os.path.exists(font_path):
            font_path = '/Library/Fonts/Arial.ttf'
        try:
            font_body = ImageFont.truetype(font_path, 22)
            font_header = ImageFont.truetype(font_path, 28)
        except Exception:
            font_body = ImageFont.load_default()
            font_header = font_body

        width, height = 1240, 1754
        margin_x, margin_y = 80, 80
        usable_height = height - (2 * margin_y) - 60

        wrapped_lines = []
        for line in lines:
            stripped = line.strip()
            if not stripped:
                wrapped_lines.append('')
                continue
            sub_lines = textwrap.wrap(stripped, width=72)
            wrapped_lines.extend(sub_lines)
            wrapped_lines.append('')

        line_height = 34
        max_lines_per_page = max(10, (usable_height - 120) // line_height)

        page_chunks = []
        curr_chunk = []
        for l in wrapped_lines:
            curr_chunk.append(l)
            if len(curr_chunk) >= max_lines_per_page:
                page_chunks.append(curr_chunk)
                curr_chunk = []
        if curr_chunk:
            page_chunks.append(curr_chunk)
        if not page_chunks:
            page_chunks = [['(Empty Document)']]

        pages = []
        total_pages = len(page_chunks)
        for p_idx, chunk in enumerate(page_chunks, 1):
            img = Image.new('RGB', (width, height), color=(255, 255, 255))
            draw = ImageDraw.Draw(img)

            draw.rectangle([margin_x, margin_y, width - margin_x, margin_y + 60], fill=(245, 247, 250))
            draw.text((margin_x + 20, margin_y + 15), title[:60], fill=(20, 20, 30), font=font_header)

            y = margin_y + 90
            for l in chunk:
                draw.text((margin_x, y), l, fill=(40, 40, 40), font=font_body)
                y += line_height

            footer_text = f"Page {p_idx} of {total_pages}"
            draw.text((width // 2 - 50, height - margin_y), footer_text, fill=(120, 120, 120), font=font_body)
            pages.append(img)

        return pages

    def extract_docx_text(docx_path):
        try:
            doc = docx.Document(docx_path)
            lines = []
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    lines.append(t)
            for table in doc.tables:
                for row in table.rows:
                    row_txt = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_txt:
                        lines.append(row_txt)
            return lines
        except Exception as e:
            app.logger.error(f"Error reading docx: {e}")
            return [f"Uploaded DOCX document: {os.path.basename(docx_path)}"]

    def allowed_file(filename):
        return '.' in filename and \
               filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

    def remove_file_safely(relative_path):
        if not relative_path:
            return
        clean_rel = relative_path.replace('static/', '', 1) if relative_path.startswith('static/') else relative_path
        abs_path = os.path.join(app.root_path, 'static', clean_rel)
        if os.path.exists(abs_path) and os.path.isfile(abs_path):
            try:
                os.remove(abs_path)
            except Exception as e:
                app.logger.warning(f"Error deleting file {abs_path}: {e}")

    @app.context_processor
    def inject_global_vault_data():
        try:
            all_papers = Paper.query.all()
            total_papers = len(all_papers)
            total_solutions = Solution.query.count()
            
            course_paper_counts = {}
            for p in all_papers:
                course_paper_counts[p.course] = course_paper_counts.get(p.course, 0) + 1
            
            courses = sorted(list(course_paper_counts.keys()))
            all_years = sorted(list({p.year for p in all_papers if p.year}), reverse=True)
            distinct_upload_years = sorted(list({p.upload_date.year for p in all_papers if p.upload_date}), reverse=True)
            
            return {
                'global_total_papers': total_papers,
                'global_total_solutions': total_solutions,
                'global_courses': courses,
                'global_course_counts': course_paper_counts,
                'global_exam_years': all_years,
                'global_upload_years': distinct_upload_years
            }
        except Exception:
            return {
                'global_total_papers': 0,
                'global_total_solutions': 0,
                'global_courses': [],
                'global_course_counts': {},
                'global_exam_years': [],
                'global_upload_years': []
            }

    @app.route('/')
    def index():
        query_course = request.args.get('course', '').strip()
        query_year = request.args.get('year', '').strip()
        query_upload_year = request.args.get('upload_year', '').strip()
        query_exam = request.args.get('exam_type', '').strip()

        paper_query = Paper.query

        if query_course:
            paper_query = paper_query.filter(Paper.course.ilike(f'%{query_course}%'))
        if query_year:
            try:
                paper_query = paper_query.filter(Paper.year == int(query_year))
            except ValueError:
                pass
        if query_upload_year:
            try:
                paper_query = paper_query.filter(db.extract('year', Paper.upload_date) == int(query_upload_year))
            except Exception:
                pass
        if query_exam:
            paper_query = paper_query.filter(Paper.exam_type == query_exam)

        papers = paper_query.order_by(Paper.upload_date.desc()).all()

        all_papers = Paper.query.all()
        distinct_upload_years = sorted(list({p.upload_date.year for p in all_papers if p.upload_date}), reverse=True)

        course_paper_counts = {}
        for p in all_papers:
            course_paper_counts[p.course] = course_paper_counts.get(p.course, 0) + 1

        all_courses = db.session.query(Paper.course).distinct().order_by(Paper.course).all()
        distinct_courses = [c[0] for c in all_courses]

        all_years = db.session.query(Paper.year).distinct().order_by(Paper.year.desc()).all()
        distinct_years = [y[0] for y in all_years]

        all_exams = db.session.query(Paper.exam_type).distinct().order_by(Paper.exam_type).all()
        distinct_exams = [e[0] for e in all_exams]

        standard_exams = ['CT1', 'CT2', 'End Sem']
        for exam in standard_exams:
            if exam not in distinct_exams:
                distinct_exams.append(exam)

        stats = {
            'total_papers': Paper.query.count(),
            'total_solutions': Solution.query.count(),
            'total_courses': len(distinct_courses)
        }

        current_year = datetime.now().year
        year_choices = list(range(current_year, 2014, -1))

        return render_template(
            'index.html',
            papers=papers,
            stats=stats,
            distinct_courses=distinct_courses,
            distinct_years=distinct_years,
            distinct_upload_years=distinct_upload_years,
            course_paper_counts=course_paper_counts,
            distinct_exams=distinct_exams,
            year_choices=year_choices,
            selected_course=query_course,
            selected_year=query_year,
            selected_upload_year=query_upload_year,
            selected_exam=query_exam
        )

    @app.route('/api/courses')
    def api_courses():
        q = request.args.get('q', '').strip()
        query = db.session.query(Paper.course).distinct()
        if q:
            query = query.filter(Paper.course.ilike(f'%{q}%'))
        courses = [c[0] for c in query.order_by(Paper.course).limit(10).all()]
        return jsonify(courses)

    @app.route('/api/check-duplicate')
    def api_check_duplicate():
        course = request.args.get('course', '').strip()
        year = request.args.get('year', '').strip()
        exam_type = request.args.get('exam_type', '').strip()

        if not (course and year and exam_type):
            return jsonify({'exists': False})

        try:
            year_int = int(year)
        except ValueError:
            return jsonify({'exists': False})

        existing = Paper.query.filter(
            Paper.course.ilike(course),
            Paper.year == year_int,
            Paper.exam_type.ilike(exam_type)
        ).first()

        if existing:
            return jsonify({
                'exists': True,
                'paper_id': existing.id,
                'course': existing.course,
                'year': existing.year,
                'exam_type': existing.exam_type,
                'page_count': existing.page_count,
                'solution_count': existing.solution_count,
                'message': f"A paper for '{existing.course}' ({existing.exam_type} {existing.year}) already exists with {existing.page_count} page(s) and {existing.solution_count} solution(s)."
            })
        
        return jsonify({'exists': False})

    @app.route('/api/papers')
    def api_papers():
        q = request.args.get('q', '').strip()
        year = request.args.get('year', '').strip()
        upload_year = request.args.get('upload_year', '').strip()
        exam = request.args.get('exam_type', '').strip()

        query = Paper.query
        if q:
            query = query.filter(Paper.course.ilike(f'%{q}%'))
        if year:
            try:
                query = query.filter(Paper.year == int(year))
            except ValueError:
                pass
        if upload_year:
            try:
                query = query.filter(db.extract('year', Paper.upload_date) == int(upload_year))
            except Exception:
                pass
        if exam:
            query = query.filter(Paper.exam_type == exam)

        papers = query.order_by(Paper.upload_date.desc()).all()
        return jsonify([p.to_dict() for p in papers])

    @app.route('/paper/<int:paper_id>')
    def paper_detail(paper_id):
        import re
        paper = db.get_or_404(Paper, paper_id)

        def clean_q_label(raw):
            r = (raw or '1').strip()
            if not r or r.lower() in ('general', 'none'):
                return '1'
            if r.lower().startswith('question'):
                r = r[8:].strip() or '1'
            elif r.lower().startswith('q') and len(r) > 1 and r[1].isdigit():
                r = r[1:].strip()
            return r

        grouped_solutions = {}
        for sol in paper.solutions:
            q = clean_q_label(sol.question_number)
            if q not in grouped_solutions:
                grouped_solutions[q] = []
            grouped_solutions[q].append(sol)

        def q_sort_key(key):
            m = re.search(r'\d+', key)
            if 'full' in key.lower() or 'all' in key.lower():
                return (2, 999999, key)
            if m:
                return (0, int(m.group()), key)
            return (1, 0, key)

        sorted_keys = sorted(grouped_solutions.keys(), key=q_sort_key)
        ordered_grouped = {k: grouped_solutions[k] for k in sorted_keys}

        answered_questions = []
        for k in sorted_keys:
            sol_list = ordered_grouped[k]
            answered_questions.append({
                'q_num': k,
                'count': len(sol_list),
                'has_multiple': len(sol_list) > 1
            })

        course_total_papers = Paper.query.filter(Paper.course.ilike(paper.course)).count()

        return render_template(
            'paper_detail.html',
            paper=paper,
            grouped_solutions=ordered_grouped,
            question_keys=sorted_keys,
            answered_questions=answered_questions,
            course_total_papers=course_total_papers
        )

    @app.route('/paper/upload', methods=['POST'])
    def upload_paper():
        course = request.form.get('course', '').strip()
        year = request.form.get('year', '').strip()
        exam_type = request.form.get('exam_type', '').strip()
        force = request.form.get('force', 'false').lower() == 'true'

        if not course or not year or not exam_type:
            flash('Please complete all required fields (Course, Year, Exam Type).', 'danger')
            return redirect(url_for('index'))

        try:
            year_int = int(year)
        except ValueError:
            flash('Invalid year provided.', 'danger')
            return redirect(url_for('index'))

        if not force:
            existing = Paper.query.filter(
                Paper.course.ilike(course),
                Paper.year == year_int,
                Paper.exam_type.ilike(exam_type)
            ).first()
            if existing:
                flash(f"Paper for {course} ({exam_type} {year_int}) already exists. Duplicate warning triggered.", 'warning')
                return redirect(url_for('paper_detail', paper_id=existing.id))

        uploaded_files = request.files.getlist('paper_files')
        valid_files = [f for f in uploaded_files if f and f.filename and allowed_file(f.filename)]

        if not valid_files:
            flash('Please select at least one valid question paper file (PDF, DOCX, TXT, PNG, JPG, WEBP).', 'danger')
            return redirect(url_for('index'))

        new_paper = Paper(
            course=course,
            year=year_int,
            exam_type=exam_type,
            upload_date=datetime.utcnow()
        )
        db.session.add(new_paper)
        db.session.flush()

        first_file = valid_files[0]
        ext = first_file.filename.rsplit('.', 1)[1].lower()

        if ext == 'pdf':
            orig_name = f"paper_{new_paper.id}_original_{uuid.uuid4().hex[:8]}.pdf"
            orig_abs_path = os.path.join(app.config['PAPERS_UPLOAD_DIR'], orig_name)
            first_file.save(orig_abs_path)
            new_paper.original_file_path = f"uploads/papers/{orig_name}"
            new_paper.original_filename = secure_filename(first_file.filename)
            new_paper.file_type = 'pdf'

            try:
                pdf = pdfium.PdfDocument(orig_abs_path)
                page_count = len(pdf)
                for idx in range(page_count):
                    page = pdf[idx]
                    pil_img = page.render(scale=2.0).to_pil()
                    page_name = f"paper_{new_paper.id}_p{idx+1}_{uuid.uuid4().hex[:8]}.png"
                    page_path = os.path.join(app.config['PAPERS_UPLOAD_DIR'], page_name)
                    pil_img.save(page_path, format='PNG')
                    paper_page = PaperPage(
                        paper_id=new_paper.id,
                        file_path=f"uploads/papers/{page_name}",
                        page_number=idx+1
                    )
                    db.session.add(paper_page)
            except Exception as e:
                app.logger.error(f"Error rendering PDF: {e}")
                flash(f"Uploaded PDF saved, but encountered an issue previewing pages: {e}", 'warning')

        elif ext in ['docx', 'doc']:
            orig_name = f"paper_{new_paper.id}_original_{uuid.uuid4().hex[:8]}.{ext}"
            orig_abs_path = os.path.join(app.config['PAPERS_UPLOAD_DIR'], orig_name)
            first_file.save(orig_abs_path)
            new_paper.original_file_path = f"uploads/papers/{orig_name}"
            new_paper.original_filename = secure_filename(first_file.filename)
            new_paper.file_type = 'docx'

            lines = extract_docx_text(orig_abs_path)
            title = f"{course} - {exam_type} ({year_int})"
            pil_pages = render_text_to_page_images(lines, title=title)
            for idx, pil_img in enumerate(pil_pages, 1):
                page_name = f"paper_{new_paper.id}_p{idx}_{uuid.uuid4().hex[:8]}.png"
                page_path = os.path.join(app.config['PAPERS_UPLOAD_DIR'], page_name)
                pil_img.save(page_path, format='PNG')
                paper_page = PaperPage(
                    paper_id=new_paper.id,
                    file_path=f"uploads/papers/{page_name}",
                    page_number=idx
                )
                db.session.add(paper_page)

        elif ext == 'txt':
            orig_name = f"paper_{new_paper.id}_original_{uuid.uuid4().hex[:8]}.txt"
            orig_abs_path = os.path.join(app.config['PAPERS_UPLOAD_DIR'], orig_name)
            first_file.save(orig_abs_path)
            new_paper.original_file_path = f"uploads/papers/{orig_name}"
            new_paper.original_filename = secure_filename(first_file.filename)
            new_paper.file_type = 'txt'

            try:
                with open(orig_abs_path, 'r', encoding='utf-8', errors='ignore') as f:
                    lines = f.readlines()
            except Exception:
                lines = ["Unable to decode text file."]

            title = f"{course} - {exam_type} ({year_int})"
            pil_pages = render_text_to_page_images(lines, title=title)
            for idx, pil_img in enumerate(pil_pages, 1):
                page_name = f"paper_{new_paper.id}_p{idx}_{uuid.uuid4().hex[:8]}.png"
                page_path = os.path.join(app.config['PAPERS_UPLOAD_DIR'], page_name)
                pil_img.save(page_path, format='PNG')
                paper_page = PaperPage(
                    paper_id=new_paper.id,
                    file_path=f"uploads/papers/{page_name}",
                    page_number=idx
                )
                db.session.add(paper_page)

        else:
            new_paper.file_type = 'images'
            for idx, file in enumerate(valid_files, start=1):
                original_ext = file.filename.rsplit('.', 1)[1].lower()
                unique_filename = f"paper_{new_paper.id}_p{idx}_{uuid.uuid4().hex[:8]}.{original_ext}"
                file_path = os.path.join(app.config['PAPERS_UPLOAD_DIR'], unique_filename)
                file.save(file_path)

                rel_path = f"uploads/papers/{unique_filename}"
                paper_page = PaperPage(
                    paper_id=new_paper.id,
                    file_path=rel_path,
                    page_number=idx
                )
                db.session.add(paper_page)

        db.session.commit()
        flash(f"Question Paper for '{course}' ({exam_type} {year_int}) uploaded successfully with {new_paper.page_count} page(s)!", 'success')
        return redirect(url_for('paper_detail', paper_id=new_paper.id))

    @app.route('/paper/<int:paper_id>/solution/upload', methods=['POST'])
    def upload_solution(paper_id):
        paper = db.get_or_404(Paper, paper_id)

        contributor = request.form.get('contributor', '').strip() or 'Community Contributor'
        notes = request.form.get('notes', '').strip()
        raw_q = request.form.get('question_number', '').strip()
        if not raw_q or raw_q.lower() in ('general', 'none'):
            question_number = '1'
        elif raw_q.lower().startswith('question'):
            question_number = raw_q[8:].strip() or '1'
        elif raw_q.lower().startswith('q') and len(raw_q) > 1 and raw_q[1].isdigit():
            question_number = raw_q[1:].strip()
        else:
            question_number = raw_q
        text_content = request.form.get('text_content', '').strip()

        uploaded_files = request.files.getlist('solution_files')
        valid_files = [
            file for file in uploaded_files
            if file and file.filename and allowed_file(file.filename)
        ]

        if not text_content and not valid_files:
            flash('Please type or paste a text answer, or upload solution files.', 'danger')
            return redirect(url_for('paper_detail', paper_id=paper.id) + '#solutions')

        sol_type = 'images'
        if text_content and valid_files:
            sol_type = 'both'
        elif text_content:
            sol_type = 'text'

        new_solution = Solution(
            paper_id=paper.id,
            contributor=contributor,
            notes=notes,
            question_number=question_number,
            solution_type=sol_type,
            text_content=text_content if text_content else None,
            upload_date=datetime.utcnow()
        )
        db.session.add(new_solution)
        db.session.flush()

        page_idx = 1
        for file in valid_files:
            original_ext = file.filename.rsplit('.', 1)[1].lower()
            if original_ext == 'pdf':
                temp_pdf_name = f"sol_{new_solution.id}_temp_{uuid.uuid4().hex[:8]}.pdf"
                temp_pdf_path = os.path.join(app.config['SOLUTIONS_UPLOAD_DIR'], temp_pdf_name)
                file.save(temp_pdf_path)
                try:
                    pdf = pdfium.PdfDocument(temp_pdf_path)
                    for p_i in range(len(pdf)):
                        page = pdf[p_i]
                        pil_img = page.render(scale=2.0).to_pil()
                        img_name = f"sol_{new_solution.id}_p{page_idx}_{uuid.uuid4().hex[:8]}.png"
                        img_path = os.path.join(app.config['SOLUTIONS_UPLOAD_DIR'], img_name)
                        pil_img.save(img_path, format='PNG')
                        db.session.add(SolutionPage(
                            solution_id=new_solution.id,
                            file_path=f"uploads/solutions/{img_name}",
                            page_number=page_idx
                        ))
                        page_idx += 1
                finally:
                    if os.path.exists(temp_pdf_path):
                        os.remove(temp_pdf_path)
            else:
                img_name = f"sol_{new_solution.id}_p{page_idx}_{uuid.uuid4().hex[:8]}.{original_ext}"
                img_path = os.path.join(app.config['SOLUTIONS_UPLOAD_DIR'], img_name)
                file.save(img_path)
                db.session.add(SolutionPage(
                    solution_id=new_solution.id,
                    file_path=f"uploads/solutions/{img_name}",
                    page_number=page_idx
                ))
                page_idx += 1

        db.session.commit()
        flash(f"Solution for Question '{question_number}' submitted successfully!", 'success')
        return redirect(url_for('paper_detail', paper_id=paper.id) + '#solutions')

    @app.route('/paper/<int:paper_id>/delete', methods=['POST'])
    def delete_paper(paper_id):
        paper = db.get_or_404(Paper, paper_id)
        course_name = paper.course

        if paper.original_file_path:
            remove_file_safely(paper.original_file_path)

        for solution in paper.solutions:
            for sp in solution.pages:
                remove_file_safely(sp.file_path)

        for pp in paper.pages:
            remove_file_safely(pp.file_path)

        db.session.delete(paper)
        db.session.commit()
        flash(f"Paper '{course_name}' and all associated pages & solutions have been permanently deleted.", 'info')
        return redirect(url_for('index'))

    @app.route('/solution/<int:solution_id>/delete', methods=['POST'])
    def delete_solution(solution_id):
        solution = db.get_or_404(Solution, solution_id)
        paper_id = solution.paper_id

        for sp in solution.pages:
            remove_file_safely(sp.file_path)

        db.session.delete(solution)
        db.session.commit()
        flash("Solution deleted successfully.", 'info')
        return redirect(url_for('paper_detail', paper_id=paper_id) + '#solutions')

    @app.route('/paper/<int:paper_id>/page/<int:page_id>/download')
    def download_paper_page(paper_id, page_id):
        paper = db.get_or_404(Paper, paper_id)
        page = PaperPage.query.filter_by(id=page_id, paper_id=paper.id).first_or_404()
        clean_rel = page.file_path.replace('static/', '', 1) if page.file_path.startswith('static/') else page.file_path
        abs_path = os.path.join(app.root_path, 'static', clean_rel)

        if not os.path.exists(abs_path):
            abort(404, description="Page file not found.")

        ext = abs_path.rsplit('.', 1)[-1]
        download_name = f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}_page_{page.page_number}.{ext}"
        return send_file(abs_path, as_attachment=True, download_name=download_name)

    @app.route('/paper/<int:paper_id>/download-original')
    def download_original_paper(paper_id):
        paper = db.get_or_404(Paper, paper_id)
        if not paper.original_file_path:
            return redirect(url_for('download_paper_pdf', paper_id=paper.id))

        clean_rel = paper.original_file_path.replace('static/', '', 1) if paper.original_file_path.startswith('static/') else paper.original_file_path
        abs_path = os.path.join(app.root_path, 'static', clean_rel)

        if not os.path.exists(abs_path):
            abort(404, description="Original document file not found.")

        ext = abs_path.rsplit('.', 1)[-1]
        download_name = paper.original_filename or f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}.{ext}"
        return send_file(abs_path, as_attachment=True, download_name=download_name)

    @app.route('/paper/<int:paper_id>/download-pdf')
    def download_paper_pdf(paper_id):
        paper = db.get_or_404(Paper, paper_id)

        if paper.original_file_path and paper.original_file_path.lower().endswith('.pdf'):
            clean_rel = paper.original_file_path.replace('static/', '', 1) if paper.original_file_path.startswith('static/') else paper.original_file_path
            abs_path = os.path.join(app.root_path, 'static', clean_rel)
            if os.path.exists(abs_path):
                download_name = f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}.pdf"
                return send_file(abs_path, as_attachment=True, download_name=download_name, mimetype='application/pdf')

        if not paper.pages:
            flash("No pages available to generate PDF.", 'warning')
            return redirect(url_for('paper_detail', paper_id=paper.id))

        pil_images = []
        for page in paper.pages:
            clean_rel = page.file_path.replace('static/', '', 1) if page.file_path.startswith('static/') else page.file_path
            abs_path = os.path.join(app.root_path, 'static', clean_rel)
            if os.path.exists(abs_path):
                img = Image.open(abs_path)
                if img.mode != 'RGB':
                    img = img.convert('RGB')
                pil_images.append(img)

        if not pil_images:
            flash("Could not load page images to compile PDF.", 'danger')
            return redirect(url_for('paper_detail', paper_id=paper.id))

        pdf_bytes = io.BytesIO()
        pil_images[0].save(
            pdf_bytes,
            format='PDF',
            save_all=True,
            append_images=pil_images[1:]
        )
        pdf_bytes.seek(0)
        download_name = f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}.pdf"
        return send_file(
            pdf_bytes,
            mimetype='application/pdf',
            as_attachment=True,
            download_name=download_name
        )

    @app.route('/paper/<int:paper_id>/download-complete-pdf')
    def download_paper_with_solutions_pdf(paper_id):
        import re
        paper = db.get_or_404(Paper, paper_id)

        all_pil_images = []

        # 1. Collect Question Paper Pages
        for page in paper.pages:
            clean_rel = page.file_path.replace('static/', '', 1) if page.file_path.startswith('static/') else page.file_path
            abs_path = os.path.join(app.root_path, 'static', clean_rel)
            if os.path.exists(abs_path):
                try:
                    img = Image.open(abs_path)
                    if img.mode != 'RGB':
                        img = img.convert('RGB')
                    all_pil_images.append(img)
                except Exception as e:
                    app.logger.warning(f"Error reading paper page {abs_path}: {e}")

        # Fallback for original PDF if no paper.pages images exist
        if not all_pil_images and paper.original_file_path and paper.original_file_path.lower().endswith('.pdf'):
            clean_rel = paper.original_file_path.replace('static/', '', 1) if paper.original_file_path.startswith('static/') else paper.original_file_path
            abs_path = os.path.join(app.root_path, 'static', clean_rel)
            if os.path.exists(abs_path):
                try:
                    import pypdfium2 as pdfium
                    pdf = pdfium.PdfDocument(abs_path)
                    for p in pdf:
                        pil_img = p.render(scale=2.0).to_pil()
                        if pil_img.mode != 'RGB':
                            pil_img = pil_img.convert('RGB')
                        all_pil_images.append(pil_img)
                except Exception as e:
                    app.logger.warning(f"Error reading pdfium pages: {e}")

        # 2. Add Solutions Section Cover & Pages if solutions exist
        if paper.solutions:
            width, height = 1240, 1754
            cover = Image.new('RGB', (width, height), color=(255, 255, 255))
            draw = ImageDraw.Draw(cover)
            
            font_path = '/System/Library/Fonts/Helvetica.ttc'
            if not os.path.exists(font_path):
                font_path = '/Library/Fonts/Arial.ttf'
            try:
                f_title = ImageFont.truetype(font_path, 42)
                f_sub = ImageFont.truetype(font_path, 26)
                f_body = ImageFont.truetype(font_path, 22)
            except Exception:
                f_title = ImageFont.load_default()
                f_sub = f_title
                f_body = f_title

            # Modern header bar
            draw.rectangle([60, 60, width - 60, 160], fill=(17, 20, 26))
            draw.text((90, 95), "PYQ VAULT — PREVIOUS YEAR QUESTIONS & SOLUTIONS", fill=(255, 255, 255), font=f_sub)

            # Title & Metadata
            draw.text((90, 260), "SOLUTIONS & VERIFIED ANSWERS", fill=(17, 20, 26), font=f_title)
            draw.line([(90, 330), (width - 90, 330)], fill=(225, 230, 240), width=3)

            draw.text((90, 360), f"Course: {paper.course}", fill=(30, 41, 59), font=f_sub)
            draw.text((90, 410), f"Examination: {paper.exam_type} ({paper.year})", fill=(71, 85, 105), font=f_sub)
            draw.text((90, 460), f"Total Verified Community Solutions: {len(paper.solutions)}", fill=(71, 85, 105), font=f_sub)

            # Question index list
            y_pos = 550
            draw.text((90, y_pos), "INDEX OF QUESTION SOLUTIONS:", fill=(15, 23, 42), font=f_sub)
            y_pos += 45
            for sol in paper.solutions:
                q_text = sol.question_number or '1'
                if q_text.lower() in ('general', 'none'):
                    q_text = '1'
                type_desc = "Typed Explanation" if sol.solution_type == 'text' else ("Handwritten Pages" if sol.solution_type == 'images' else "Typed + Handwritten")
                draw.text((110, y_pos), f"• Question {q_text}: Answer by {sol.contributor or 'Scholar'} ({type_desc})", fill=(51, 65, 85), font=f_body)
                y_pos += 36
                if y_pos > height - 180:
                    break

            draw.text((90, height - 90), "Generated by PYQ Vault — Past Papers & Solutions Archive", fill=(148, 163, 184), font=f_body)
            all_pil_images.append(cover)

            # Sort solutions by question number so they appear in order
            def sol_sort_key(s):
                m = re.search(r'\d+', s.question_number or '1')
                if m:
                    return (0, int(m.group()))
                return (1, 0)
            
            sorted_sols = sorted(paper.solutions, key=sol_sort_key)

            for sol in sorted_sols:
                q_label = sol.question_number or '1'
                if q_label.lower() in ('general', 'none'):
                    q_label = '1'
                sol_title = f"{paper.course} | Question {q_label} Solution (by {sol.contributor or 'Scholar'})"

                # Typed Text Solution
                if sol.text_content and sol.text_content.strip():
                    text_lines = []
                    if sol.notes:
                        text_lines.append(f"Notes: {sol.notes}")
                        text_lines.append("-" * 40)
                    text_lines.extend(sol.text_content.splitlines())
                    text_pages = render_text_to_page_images(text_lines, title=sol_title)
                    all_pil_images.extend(text_pages)

                # Handwritten Image Pages
                for sp in sol.pages:
                    clean_rel = sp.file_path.replace('static/', '', 1) if sp.file_path.startswith('static/') else sp.file_path
                    abs_path = os.path.join(app.root_path, 'static', clean_rel)
                    if os.path.exists(abs_path):
                        try:
                            img = Image.open(abs_path)
                            if img.mode != 'RGB':
                                img = img.convert('RGB')
                            all_pil_images.append(img)
                        except Exception as e:
                            app.logger.warning(f"Error reading solution page {abs_path}: {e}")

        if not all_pil_images:
            flash("No pages available to generate compiled PDF.", 'warning')
            return redirect(url_for('paper_detail', paper_id=paper.id))

        pdf_bytes = io.BytesIO()
        all_pil_images[0].save(
            pdf_bytes,
            format='PDF',
            save_all=True,
            append_images=all_pil_images[1:]
        )
        pdf_bytes.seek(0)
        download_name = f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}_Paper_and_Solutions.pdf"
        return send_file(
            pdf_bytes,
            mimetype='application/pdf',
            as_attachment=True,
            download_name=download_name
        )

    @app.route('/paper/<int:paper_id>/download-zip')
    def download_paper_zip(paper_id):
        paper = db.get_or_404(Paper, paper_id)
        if not paper.pages:
            flash("No pages available in this paper to download.", 'warning')
            return redirect(url_for('paper_detail', paper_id=paper.id))

        memory_file = io.BytesIO()
        with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            if paper.original_file_path:
                clean_rel = paper.original_file_path.replace('static/', '', 1) if paper.original_file_path.startswith('static/') else paper.original_file_path
                orig_abs = os.path.join(app.root_path, 'static', clean_rel)
                if os.path.exists(orig_abs):
                    zf.write(orig_abs, arcname=paper.original_filename or os.path.basename(orig_abs))

            for page in paper.pages:
                clean_rel = page.file_path.replace('static/', '', 1) if page.file_path.startswith('static/') else page.file_path
                abs_path = os.path.join(app.root_path, 'static', clean_rel)
                if os.path.exists(abs_path):
                    ext = abs_path.rsplit('.', 1)[-1]
                    archive_name = f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}_page_{page.page_number:02d}.{ext}"
                    zf.write(abs_path, arcname=archive_name)

        memory_file.seek(0)
        zip_filename = f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}_full_paper.zip"
        return send_file(
            memory_file,
            mimetype='application/zip',
            as_attachment=True,
            download_name=zip_filename
        )

    @app.route('/course/<path:course_name>/download-all')
    def download_course_papers_zip(course_name):
        clean_course = unquote(course_name).strip()
        papers = Paper.query.filter(Paper.course.ilike(clean_course)).order_by(Paper.year.desc(), Paper.exam_type).all()

        if not papers:
            flash(f"No papers found for course '{clean_course}'.", 'warning')
            return redirect(url_for('index'))

        safe_course = secure_filename(clean_course) or "course_papers"
        memory_file = io.BytesIO()

        with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            manifest_lines = [
                f"PYQ Vault Course Archive: {clean_course}",
                f"Generated at: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
                f"Total Question Papers: {len(papers)}",
                "=" * 50,
                ""
            ]

            for paper in papers:
                folder_name = f"{paper.year}_{secure_filename(paper.exam_type)}"
                manifest_lines.append(f"Paper: {paper.course} ({paper.exam_type} {paper.year})")
                manifest_lines.append(f"  - Pages: {paper.page_count}")
                manifest_lines.append(f"  - Solutions: {paper.solution_count}")
                manifest_lines.append("")

                if paper.original_file_path:
                    clean_rel = paper.original_file_path.replace('static/', '', 1) if paper.original_file_path.startswith('static/') else paper.original_file_path
                    orig_abs = os.path.join(app.root_path, 'static', clean_rel)
                    if os.path.exists(orig_abs):
                        orig_ext = orig_abs.rsplit('.', 1)[-1]
                        arc_name = f"{safe_course}/{folder_name}/{paper.original_filename or f'paper_{paper.id}_original.{orig_ext}'}"
                        zf.write(orig_abs, arcname=arc_name)

                for page in paper.pages:
                    clean_rel = page.file_path.replace('static/', '', 1) if page.file_path.startswith('static/') else page.file_path
                    abs_path = os.path.join(app.root_path, 'static', clean_rel)
                    if os.path.exists(abs_path):
                        ext = abs_path.rsplit('.', 1)[-1]
                        arc_name = f"{safe_course}/{folder_name}/pages/{safe_course}_{paper.exam_type}_{paper.year}_page_{page.page_number:02d}.{ext}"
                        zf.write(abs_path, arcname=arc_name)

                for sol in paper.solutions:
                    q_slug = secure_filename(sol.question_number or 'general')
                    if sol.text_content:
                        txt_arc = f"{safe_course}/{folder_name}/solutions/Q_{q_slug}_sol_{sol.id}.txt"
                        zf.writestr(txt_arc, f"Course: {paper.course}\nExam: {paper.exam_type} ({paper.year})\nQuestion: {sol.question_number}\nContributor: {sol.contributor}\nNotes: {sol.notes or 'None'}\n\nSolution Content:\n{sol.text_content}\n")

                    for sp in sol.pages:
                        clean_rel = sp.file_path.replace('static/', '', 1) if sp.file_path.startswith('static/') else sp.file_path
                        sp_abs = os.path.join(app.root_path, 'static', clean_rel)
                        if os.path.exists(sp_abs):
                            ext = sp_abs.rsplit('.', 1)[-1]
                            sp_arc = f"{safe_course}/{folder_name}/solutions/Q_{q_slug}_sol_{sol.id}_page_{sp.page_number:02d}.{ext}"
                            zf.write(sp_abs, arcname=sp_arc)

            zf.writestr(f"{safe_course}/README.txt", "\n".join(manifest_lines))

        memory_file.seek(0)
        zip_filename = f"{safe_course}_all_years_papers.zip"
        return send_file(
            memory_file,
            mimetype='application/zip',
            as_attachment=True,
            download_name=zip_filename
        )

    @app.route('/solution/<int:solution_id>/page/<int:page_id>/download')
    def download_solution_page(solution_id, page_id):
        solution = db.get_or_404(Solution, solution_id)
        page = SolutionPage.query.filter_by(id=page_id, solution_id=solution.id).first_or_404()
        clean_rel = page.file_path.replace('static/', '', 1) if page.file_path.startswith('static/') else page.file_path
        abs_path = os.path.join(app.root_path, 'static', clean_rel)

        if not os.path.exists(abs_path):
            abort(404, description="Solution page file not found.")

        ext = abs_path.rsplit('.', 1)[-1]
        paper = solution.paper
        download_name = f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}_solution_{solution.id}_p{page.page_number}.{ext}"
        return send_file(abs_path, as_attachment=True, download_name=download_name)

    @app.route('/solution/<int:solution_id>/download-zip')
    def download_solution_zip(solution_id):
        solution = db.get_or_404(Solution, solution_id)
        paper = solution.paper
        if not solution.pages:
            flash("No pages available for this solution.", 'warning')
            return redirect(url_for('paper_detail', paper_id=paper.id))

        memory_file = io.BytesIO()
        with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            for page in solution.pages:
                clean_rel = page.file_path.replace('static/', '', 1) if page.file_path.startswith('static/') else page.file_path
                abs_path = os.path.join(app.root_path, 'static', clean_rel)
                if os.path.exists(abs_path):
                    ext = abs_path.rsplit('.', 1)[-1]
                    archive_name = f"solution_{solution.id}_page_{page.page_number:02d}.{ext}"
                    zf.write(abs_path, arcname=archive_name)

        memory_file.seek(0)
        zip_filename = f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}_solution_{solution.id}.zip"
        return send_file(
            memory_file,
            mimetype='application/zip',
            as_attachment=True,
            download_name=zip_filename
        )

    @app.route('/solution/<int:solution_id>/download-text')
    def download_solution_text(solution_id):
        solution = db.get_or_404(Solution, solution_id)
        if not solution.text_content:
            flash("This solution does not contain text content to download.", 'warning')
            return redirect(url_for('paper_detail', paper_id=solution.paper_id) + '#solutions')

        paper = solution.paper
        filename = f"{secure_filename(paper.course)}_{paper.exam_type}_{paper.year}_Q{secure_filename(solution.question_number or 'general')}_sol_{solution.id}.txt"
        content = (
            f"Course: {paper.course}\n"
            f"Exam: {paper.exam_type} ({paper.year})\n"
            f"Question: {solution.question_number or 'General'}\n"
            f"Contributor: {solution.contributor}\n"
            f"Date: {solution.upload_date.strftime('%Y-%m-%d') if solution.upload_date else ''}\n"
            f"Notes: {solution.notes or 'None'}\n\n"
            f"--- SOLUTION CONTENT ---\n\n"
            f"{solution.text_content}\n"
        )
        return send_file(
            io.BytesIO(content.encode('utf-8')),
            mimetype='text/plain',
            as_attachment=True,
            download_name=filename
        )

    @app.errorhandler(404)
    def page_not_found(e):
        return render_template('404.html'), 404

    return app

app = create_app()

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5002))
    print(f"Starting PYQ Vault server on http://127.0.0.1:{port} (host: 0.0.0.0)")
    app.run(debug=True, host='0.0.0.0', port=port)
